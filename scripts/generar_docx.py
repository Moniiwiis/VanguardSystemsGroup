# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Margenes
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

ROJO_PRINCIPAL = RGBColor(196, 18, 48)  # #C41230
GRIS_OSCURO = RGBColor(36, 16, 22)
GRIS_TEXTO = RGBColor(71, 85, 105)

# Portada / Encabezado
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_inst = title_p.add_run("INACAP · PROGRAMACIÓN BACK END (TI3041)\n")
run_inst.font.name = "Arial"
run_inst.font.size = Pt(11)
run_inst.font.bold = True
run_inst.font.color.rgb = ROJO_PRINCIPAL

run_title = title_p.add_run("PRE-EVALUACIÓN SUMATIVA III (40 PUNTOS)\n")
run_title.font.name = "Arial"
run_title.font.size = Pt(18)
run_title.font.bold = True
run_title.font.color.rgb = ROJO_PRINCIPAL

run_sub = title_p.add_run("Hito 1: Validación Individual de CRUD Django + Integración Clase 10\nProyecto: Vanguard Systems Group — Ilustre Municipalidad de La Serena")
run_sub.font.name = "Arial"
run_sub.font.size = Pt(10)
run_sub.font.italic = True
run_sub.font.color.rgb = GRIS_TEXTO

doc.add_paragraph()

def add_custom_heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = ROJO_PRINCIPAL
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(139, 29, 25)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = GRIS_OSCURO
    return p

# 1. Resumen
add_custom_heading("1. Resumen de Cumplimiento Técnico y Clase 10", 1)
p_res = doc.add_paragraph()
p_res.add_run("El proyecto VanguardSystemsGroup cumple estrictamente con los cuatro criterios de la pauta de evaluación y con todas las directrices de la Clase 10:\n")

items_res = [
    ("Persistencia de Paginación en Sesión: ", "Implementada en la función _get_page_size(). Valida únicamente 5, 15 o 30 registros. Si se recibe por GET se guarda en request.session['page_size']; al refrescar o navegar sin el parámetro, se recupera el valor de la sesión."),
    ("Conservación de Filtros al Paginar: ", "Implementada en _paginate_queryset() y paginator.html. Mediante request.GET.copy() y params.pop('page', None), genera query_params = params.urlencode(). Los enlaces a páginas previas, siguientes o números nunca pierden los filtros ni el término de búsqueda."),
    ("Secuencia Segura de Consulta (5 Capas): ", "Scoping de seguridad territorial -> Búsqueda libre con Q e icontains -> Filtros exactos por GET -> Ordenamiento con Whitelist (ORDERING_MAP) -> Paginación final."),
    ("Modales para Create y Update: ", "La creación y edición de registros se realiza mediante capas modales con backdrop y formularios POST con token CSRF, actualizados dinámicamente con JavaScript."),
    ("Eliminación con SweetAlert2: ", "La función confirmDelete() intercepta el submit del formulario y despliega un cuadro de diálogo interactivo Swal.fire(), enviando el POST solo si el usuario confirma."),
    ("Pruebas Automatizadas: ", "13 tests unitarios aprobados al 100% (python manage.py test en verde).")
]
for title, desc in items_res:
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(title)
    r1.font.bold = True
    p.add_run(desc)

# 2. Inventario de CRUDs
add_custom_heading("2. Inventario Completo de Todos los CRUDs del Proyecto", 1)
doc.add_paragraph("Cada integrante del grupo puede elegir y defender uno de los siguientes 8 CRUDs completos:")

cruds_data = [
    ("Vecinos (Catastro)", "Vecino (activities)", "gestion_vecinos", "vecinos.html", "Recomendado: Scoping territorial, 6 campos con Q/tildes"),
    ("Tipo de Gestión", "TipoGestion (activities)", "gestion_tipo_gestion", "tipo_gestion.html", "Directo y limpio, ideal para defensa clara"),
    ("Tipo de Atención", "TipoAtencion (configuration)", "gestion_tipo_atencion", "tipo_atencion.html", "Incluye integridad referencial contra SubAtencion"),
    ("Sub Atención", "SubAtencion (configuration)", "gestion_sub_atencion", "sub_atencion.html", "Relación ForeignKey con select_related()"),
    ("Delegaciones", "Delegacion (accounts)", "gestion_delegaciones", "delegaciones.html", "Entidad base territorial del sistema"),
    ("Roles de Usuario", "Rol (accounts)", "gestion_roles", "roles.html", "Maneja permisos y alcance (Delegación/Central)"),
    ("Metas / Medición", "ItemMedicion (configuration)", "gestion_metas", "metas.html", "Indicadores del SGR municipal"),
    ("Usuarios / Personal", "User + Funcionario (accounts)", "gestion_usuarios", "usuarios.html", "Complejo: vincula User de Django con Funcionario")
]

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True
hdr_cells = table.rows[0].cells
hdr_titles = ["Módulo CRUD", "Modelo ORM", "Vista (views.py)", "Template", "Particularidad"]
for i, title in enumerate(hdr_titles):
    hdr_cells[i].text = title
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True
    hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9)

for crud in cruds_data:
    row_cells = table.add_row().cells
    for i, val in enumerate(crud):
        row_cells[i].text = val
        row_cells[i].paragraphs[0].runs[0].font.size = Pt(8.5)

doc.add_paragraph()

# 3. Guia de Defensa Paso a Paso
add_custom_heading("3. Guía de Defensa Individual Paso a Paso (40 Puntos)", 1)

# Criterio 1
add_custom_heading("Criterio 1: Ejecución local y login funcional (5 Pts)", 2)
p1 = doc.add_paragraph()
p1.add_run("Demostración en vivo: ").font.bold = True
p1.add_run("Ejecutar en consola 'python manage.py runserver', abrir http://127.0.0.1:8000/login/ e ingresar credenciales.\n")
p1.add_run("Código a explicar: ").font.bold = True
p1.add_run("Vista login_view en core/views.py. Explicar authenticate() y login(request, user). Las contraseñas están protegidas con hash PBKDF2/SHA256. Las vistas privadas están aseguradas con el decorador @login_required.")

# Criterio 2
add_custom_heading("Criterio 2: CRUD completo individual (15 Pts)", 2)
p2 = doc.add_paragraph()
p2.add_run("Demostración en vivo: ").font.bold = True
p2.add_run("Ejecutar las 4 operaciones (Listar en tabla, Crear en modal, Editar en modal y Eliminar con confirmación).\n")
p2.add_run("Código a explicar: ").font.bold = True
p2.add_run("Mapeo en config/urls.py -> Vista en core/views.py -> Modelo ORM -> Template:\n")
p2_items = [
    ("Listar: ", "Consulta ORM Model.objects.filter(...) y renderizado en el template."),
    ("Crear: ", "Bloque if request.method == 'POST' en la vista principal; valida campos y ejecuta Model.objects.create(...) insertando en SQLite."),
    ("Editar: ", "Vista de edición recibe el ID del registro; busca con get_object_or_404(), actualiza campos desde request.POST y ejecuta objeto.save() (UPDATE en base de datos)."),
    ("Eliminar: ", "Vista de eliminación recibe ID por POST, busca el objeto y ejecuta objeto.delete() (DELETE en base de datos).")
]
for t, d in p2_items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(t).font.bold = True
    p.add_run(d)

# Criterio 3
add_custom_heading("Criterio 3: Modal Bootstrap / UI + SweetAlert2 (8 Pts)", 2)
p3 = doc.add_paragraph()
p3.add_run("Modales Crear y Editar: ").font.bold = True
p3.add_run("Implementados en el template del CRUD con formularios POST y {% csrf_token %}. Para editar, una función JavaScript llena los inputs del modal con los datos de la fila y actualiza el atributo action del form sin recargar la página.\n")
p3.add_run("SweetAlert2 en Delete: ").font.bold = True
p3.add_run("En templates/base.html, la función confirmDelete(form, itemName) intercepta el onsubmit. Si el usuario confirma en Swal.fire(), se llama a form.submit(), ejecutando la eliminación de forma segura y visual.")

# Criterio 4
add_custom_heading("Criterio 4: Buscador + Filtros + Orden + Paginación con Sesión (12 Pts)", 2)
p4 = doc.add_paragraph()
p4.add_run("Este criterio integra todos los aprendizajes de la Clase 10:\n")
c4_items = [
    ("1. Scoping Territorial Primero: ", "Antes de aplicar búsqueda o filtros, se restringe el QuerySet según el usuario (ej: qs = qs.filter(delegacion=del_activa)). Un usuario de la Delegación Norte nunca puede 'descubrir' datos de la Delegación Sur."),
    ("2. Buscador con Q e icontains: ", "from django.db.models import Q. Se usa Q(campo1__icontains=q) | Q(campo2__icontains=q). icontains es insensible a mayúsculas y el operador pipe (|) aplica lógica OR entre varios campos."),
    ("3. Filtros Concretos: ", "Parámetros GET específicos (ej. ?estado=Activo) aplicados con filter(estado__iexact=...)."),
    ("4. Ordenamiento Seguro con Whitelist: ", "Función _apply_list_ordering() valida el parámetro 'orden' contra una lista permitida (ORDERING_MAP). 'GET propone, Django valida'."),
    ("5. Paginación y Sesión (5, 15, 30): ", "Función _get_page_size() guarda la preferencia en request.session['page_size']. Al refrescar (F5) o navegar, la sesión recuerda el tamaño seleccionado."),
    ("6. Conservar Parámetros al Paginar: ", "request.GET.copy() retira 'page' y genera query_params = params.urlencode(). Los enlaces del paginador mantienen la consulta completa (?q=...&estado=...&orden=...&page=2).")
]
for t, d in c4_items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(t).font.bold = True
    p.add_run(d)

# 4. Ticket de Salida
add_custom_heading("4. Ticket de Salida: Preguntas Típicas del Docente y Respuestas", 1)

faqs = [
    ("¿Por qué el scoping debe aplicarse antes del buscador?",
     "Por seguridad. El usuario solo debe acceder al subconjunto de datos para el cual tiene autorización. Si la búsqueda se hiciera antes del scoping, el buscador podría exponer registros confidenciales de otras delegaciones u organizaciones."),
    ("¿Qué problema resuelve la whitelist de ordenamiento (ORDERING_MAP)?",
     "Evita que el usuario inyecte nombres de campos arbitrarios o maliciosos por el parámetro GET que provoquen excepciones (FieldError) o expongan información confidencial. Regla: GET propone, Django valida."),
    ("¿Por qué page_size encaja mejor en session que en la base de datos?",
     "Porque page_size es una preferencia temporal de navegación e interfaz del usuario durante su sesión de trabajo, no un dato permanente de las reglas de negocio del dominio (como un RUT o una Delegación). Almacenarlo en sesión evita sobrecargar la base de datos."),
    ("¿Por qué usamos el objeto Q en lugar de encadenar varios .filter()?",
     "Porque al encadenar .filter() las condiciones se combinan con AND (debe cumplirse todo al mismo tiempo). El objeto Q nos permite usar el operador pipe (|) para expresar lógica OR, permitiendo que el término buscado coincida con el nombre O con el RUT O con la dirección.")
]

for q, a in faqs:
    p_q = doc.add_paragraph()
    r_q = p_q.add_run("Pregunta: " + q)
    r_q.font.bold = True
    r_q.font.color.rgb = RGBColor(139, 29, 25)
    p_a = doc.add_paragraph()
    p_a.add_run("Respuesta: ").font.bold = True
    p_a.add_run(a)

doc.save("GUIA_PREEVALUACION_CRUD_DJANGO.docx")
print("DOCX successfully generated: GUIA_PREEVALUACION_CRUD_DJANGO.docx")
